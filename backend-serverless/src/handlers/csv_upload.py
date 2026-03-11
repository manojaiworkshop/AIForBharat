"""
csv_upload.py — CSV/Excel upload to S3 per-user; optional CREATE TABLE in superadmin's PostgreSQL.

Routes (require Authorization: Bearer <JWT>):
  POST   /csv/upload                          → upload file to S3, save metadata to DynamoDB
  POST   /csv/tables/{table_id}/create-table  → create PG table from the uploaded file (superadmin DB)
  GET    /csv/tables                          → list all files + created table info
  GET    /csv/tables/{table_id}               → fetch rows (from PG if table created, else parsed from S3)
  DELETE /csv/tables/{table_id}               → delete S3 file, PG table (if exists), DynamoDB record
  OPTIONS /csv/{proxy+}                       → CORS pre-flight
"""

import json
import os
import csv
import io
import re
import ssl as ssl_mod
import uuid
import base64
import hashlib
import hmac
import logging
from datetime import datetime, timezone

import boto3
from boto3.dynamodb.conditions import Key
import pg8000.native as pg8000

logger = logging.getLogger()
logger.setLevel(logging.INFO)

CSV_TABLES_TABLE     = os.environ.get("CSV_TABLES_TABLE",     "mercury-grid-csv-tables")
DB_CONNECTIONS_TABLE = os.environ.get("DB_CONNECTIONS_TABLE", "mercury-grid-db-connections")
CSV_BUCKET           = os.environ.get("CSV_BUCKET",           "mercury-grid-csv-files")
JWT_SECRET           = os.environ.get("JWT_SECRET",           "mercury-grid-secret-change-in-prod")
ADMIN_USER_ID        = "__admin__"
MAX_ROWS_STORED      = 5000

dynamodb  = boto3.resource("dynamodb")
s3        = boto3.client("s3")
meta_tbl  = dynamodb.Table(CSV_TABLES_TABLE)
conn_tbl  = dynamodb.Table(DB_CONNECTIONS_TABLE)

CORS_HEADERS = {
    "Access-Control-Allow-Origin":  "*",
    "Access-Control-Allow-Headers": "Content-Type,Authorization",
    "Access-Control-Allow-Methods": "GET,POST,PUT,DELETE,OPTIONS",
    "Content-Type": "application/json",
}


# ── Auth helpers ──────────────────────────────────────────────

def _verify_jwt(token: str) -> dict | None:
    try:
        parts = token.split(".")
        if len(parts) != 3:
            return None
        header_b64, payload_b64, sig_b64 = parts
        msg = f"{header_b64}.{payload_b64}".encode()
        expected = base64.urlsafe_b64encode(
            hmac.new(JWT_SECRET.encode(), msg, hashlib.sha256).digest()
        ).rstrip(b"=").decode()
        if not hmac.compare_digest(sig_b64, expected):
            return None
        pad = 4 - len(payload_b64) % 4
        return json.loads(base64.urlsafe_b64decode(payload_b64 + "=" * pad))
    except Exception:
        return None


def _get_user(event: dict) -> tuple[dict | None, dict | None]:
    auth = (event.get("headers") or {}).get("Authorization") or ""
    if not auth.startswith("Bearer "):
        return None, _error(401, "Authorization header required")
    payload = _verify_jwt(auth[7:])
    if not payload:
        return None, _error(401, "Invalid or expired token")
    return payload, None


# ── Response helpers ──────────────────────────────────────────

def _ok(body: dict, status: int = 200) -> dict:
    return {"statusCode": status, "headers": CORS_HEADERS, "body": json.dumps(body, default=str)}


def _error(status: int, message: str) -> dict:
    return {"statusCode": status, "headers": CORS_HEADERS, "body": json.dumps({"error": message})}


# ── PostgreSQL helpers ────────────────────────────────────────

def _pg_connect(cfg: dict):
    kwargs = {
        "host":     cfg["host"],
        "port":     int(cfg.get("port", 5432)),
        "database": cfg["database"],
        "user":     cfg["username"],
        "password": cfg["password"],
        "timeout":  20,
    }
    if cfg.get("ssl", False):
        ctx = ssl_mod.create_default_context()
        ctx.check_hostname = False
        ctx.verify_mode    = ssl_mod.CERT_NONE
        kwargs["ssl_context"] = ctx
    else:
        kwargs["ssl_context"] = None
    return pg8000.Connection(**kwargs)


def _get_superadmin_pg_connection() -> dict | None:
    """Return the default (most recent) postgres connection saved by superadmin."""
    try:
        result = conn_tbl.query(KeyConditionExpression=Key("user_id").eq(ADMIN_USER_ID))
        items  = [i for i in result.get("Items", []) if i.get("type") == "postgres"]
        if not items:
            return None
        items.sort(key=lambda x: x.get("created_at", ""), reverse=True)
        return items[0]
    except Exception as e:
        logger.exception("_get_superadmin_pg_connection error: %s", e)
        return None


def _safe_pg_name(name: str) -> str:
    safe = re.sub(r"[^a-z0-9_]", "_", name.lower())
    safe = re.sub(r"_+", "_", safe).strip("_") or "table"
    return safe[:50]


def _safe_col(name: str) -> str:
    safe = re.sub(r"[^a-z0-9_]", "_", name.lower())
    safe = re.sub(r"_+", "_", safe).strip("_") or "col"
    return safe[:60]


# ── Parsing helpers ───────────────────────────────────────────

def _dedup_cols(cols: list) -> list:
    seen: dict = {}
    out = []
    for c in cols:
        c = c or "column"
        if c in seen:
            seen[c] += 1
            out.append(f"{c}_{seen[c]}")
        else:
            seen[c] = 0
            out.append(c)
    return out


def _parse_csv_bytes(data: bytes) -> tuple:
    text     = data.decode("utf-8-sig", errors="replace")
    reader   = csv.reader(io.StringIO(text))
    rows_raw = list(reader)
    if not rows_raw:
        return [], []
    columns = _dedup_cols([c.strip().replace(" ", "_").replace("-", "_") or "column" for c in rows_raw[0]])
    rows    = [list(r) for r in rows_raw[1:MAX_ROWS_STORED + 1]]
    return columns, rows


def _parse_excel_bytes(data: bytes) -> tuple:
    import openpyxl
    wb       = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws       = wb.active
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not all_rows:
        return [], []
    def clean(c):
        return (str(c) if c is not None else "").strip().replace(" ", "_").replace("-", "_") or "column"
    columns = _dedup_cols([clean(c) for c in all_rows[0]])
    rows    = [[str(v) if v is not None else "" for v in r] for r in all_rows[1:MAX_ROWS_STORED + 1]]
    return columns, rows


def _parse_file(raw_bytes: bytes, ext: str) -> tuple:
    if ext in ("xlsx", "xls", "xlsm"):
        return _parse_excel_bytes(raw_bytes)
    return _parse_csv_bytes(raw_bytes)


# ── Route handlers ────────────────────────────────────────────

def upload_file(event: dict) -> dict:
    """POST /csv/upload — upload raw file to S3, save metadata to DynamoDB."""
    user, err = _get_user(event)
    if err:
        return err
    user_id = user.get("sub", "anonymous")

    try:
        body = json.loads(event.get("body") or "{}")
    except json.JSONDecodeError:
        return _error(400, "Invalid JSON body")

    file_name    = (body.get("file_name") or "").strip()
    file_content = body.get("file_content") or ""
    file_type    = (body.get("file_type") or "").lower()

    if not file_name:
        return _error(400, "file_name is required")
    if not file_content:
        return _error(400, "file_content is required (base64)")

    try:
        raw_bytes = base64.b64decode(file_content)
    except Exception:
        return _error(400, "file_content must be valid base64")

    # Parse to get column info + row count
    try:
        ext = file_type or os.path.splitext(file_name)[1].lstrip(".").lower()
        columns, rows = _parse_file(raw_bytes, ext)
    except Exception as exc:
        logger.exception("Parse error")
        return _error(422, f"Failed to parse file: {exc}")

    if not columns:
        return _error(422, "File is empty or has no header row")

    safe_cols = [_safe_col(c) for c in columns]
    table_id  = str(uuid.uuid4())
    now       = datetime.now(timezone.utc).isoformat()
    row_count = len(rows)
    truncated = row_count >= MAX_ROWS_STORED
    s3_key    = f"csv/{user_id}/{table_id}/{file_name}"
    table_name = _safe_pg_name(os.path.splitext(file_name)[0])

    # Upload raw file to S3
    content_type = "text/csv" if ext == "csv" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    try:
        s3.put_object(
            Bucket=CSV_BUCKET,
            Key=s3_key,
            Body=raw_bytes,
            ContentType=content_type,
            Metadata={
                "user_id":   user_id,
                "table_id":  table_id,
                "file_name": file_name,
            },
        )
    except Exception as exc:
        logger.exception("S3 upload error")
        return _error(500, f"S3 upload failed: {exc}")

    # Save metadata to DynamoDB (no row data stored here)
    try:
        meta_tbl.put_item(Item={
            "user_id":       user_id,
            "table_id":      table_id,
            "table_name":    table_name,
            "file_name":     file_name,
            "file_ext":      ext,
            "s3_key":        s3_key,
            "col_names":     safe_cols,
            "row_count":     row_count,
            "truncated":     truncated,
            "table_created": False,
            "pg_table_name": None,
            "created_at":    now,
        })
    except Exception as exc:
        logger.exception("DynamoDB metadata error")
        return _error(500, f"Metadata save failed: {exc}")

    return _ok({
        "table_id":      table_id,
        "table_name":    table_name,
        "file_name":     file_name,
        "columns":       safe_cols,
        "row_count":     row_count,
        "truncated":     truncated,
        "table_created": False,
        "created_at":    now,
    }, 201)


def create_table(event: dict, table_id: str) -> dict:
    """POST /csv/tables/{table_id}/create-table — read from S3, create PG table in superadmin's DB."""
    user, err = _get_user(event)
    if err:
        return err
    user_id = user.get("sub", "anonymous")

    # Get metadata
    try:
        result = meta_tbl.get_item(Key={"user_id": user_id, "table_id": table_id})
        item   = result.get("Item")
    except Exception as exc:
        return _error(500, str(exc))

    if not item:
        return _error(404, "File not found")

    if item.get("table_created"):
        return _ok({
            "message":     "Table already created",
            "pg_table_name": item.get("pg_table_name"),
            "table_created": True,
            "columns":     item.get("col_names", []),
        })

    # Get superadmin's PostgreSQL connection
    db_cfg = _get_superadmin_pg_connection()
    if not db_cfg:
        return _error(400, "No PostgreSQL connection configured by admin. Ask admin to save a connection in the Admin Panel.")

    # Read file from S3
    s3_key = item.get("s3_key", "")
    try:
        obj       = s3.get_object(Bucket=CSV_BUCKET, Key=s3_key)
        raw_bytes = obj["Body"].read()
    except Exception as exc:
        return _error(500, f"Could not read file from S3: {exc}")

    # Parse the file
    ext = item.get("file_ext", "csv")
    try:
        columns, rows = _parse_file(raw_bytes, ext)
    except Exception as exc:
        return _error(422, f"Failed to parse file: {exc}")

    safe_cols = item.get("col_names") or [_safe_col(c) for c in columns]

    # Generate unique PG table name
    base_name  = item.get("table_name", "table")
    suffix     = table_id.replace("-", "")[:6]
    pg_table   = f"{base_name}_{suffix}"

    # Connect to superadmin's PostgreSQL and create table
    try:
        conn = _pg_connect(db_cfg)
    except Exception as exc:
        return _error(502, f"Could not connect to database: {exc}")

    try:
        col_defs = ", ".join(f'"{c}" TEXT' for c in safe_cols)
        conn.run(f'CREATE TABLE IF NOT EXISTS "{pg_table}" ({col_defs})')

        # Batch inserts of 200 rows
        for i in range(0, len(rows), 200):
            batch  = rows[i:i + 200]
            params = {}
            value_rows = []
            idx = 0
            for row in batch:
                placeholders = []
                for c_i in range(len(safe_cols)):
                    key = f"p{idx}"
                    params[key] = row[c_i] if c_i < len(row) else ""
                    placeholders.append(f":{key}")
                    idx += 1
                value_rows.append("(" + ", ".join(placeholders) + ")")
            col_names = ", ".join(f'"{c}"' for c in safe_cols)
            conn.run(
                f'INSERT INTO "{pg_table}" ({col_names}) VALUES {", ".join(value_rows)}',
                **params,
            )
        conn.close()
    except Exception as exc:
        logger.exception("PG write error")
        try:
            conn.close()
        except Exception:
            pass
        return _error(500, f"Database write failed: {exc}")

    # Update DynamoDB metadata
    try:
        meta_tbl.update_item(
            Key={"user_id": user_id, "table_id": table_id},
            UpdateExpression="SET table_created = :tc, pg_table_name = :pt, connection_id = :ci, connection_name = :cn",
            ExpressionAttributeValues={
                ":tc": True,
                ":pt": pg_table,
                ":ci": db_cfg.get("connection_id", ""),
                ":cn": db_cfg.get("name", ""),
            },
        )
    except Exception as exc:
        logger.exception("DynamoDB update error after PG table creation")
        return _error(500, f"Metadata update failed: {exc}")

    return _ok({
        "success":         True,
        "pg_table_name":   pg_table,
        "table_created":   True,
        "columns":         safe_cols,
        "row_count":       len(rows),
        "connection_name": db_cfg.get("name", ""),
    })


def list_tables(event: dict) -> dict:
    """GET /csv/tables — list all uploaded files for the user."""
    user, err = _get_user(event)
    if err:
        return err
    user_id = user.get("sub", "anonymous")

    try:
        result = meta_tbl.query(KeyConditionExpression=Key("user_id").eq(user_id))
        tables = result.get("Items", [])
        tables.sort(key=lambda x: x.get("created_at", ""), reverse=True)
        for t in tables:
            t["columns"] = t.pop("col_names", [])
        return _ok({"tables": tables})
    except Exception as exc:
        logger.exception("list_tables error")
        return _error(500, str(exc))


def get_table(event: dict, table_id: str) -> dict:
    """GET /csv/tables/{table_id} — fetch rows from PG (if table created) or parse directly from S3."""
    user, err = _get_user(event)
    if err:
        return err
    user_id = user.get("sub", "anonymous")

    try:
        result = meta_tbl.get_item(Key={"user_id": user_id, "table_id": table_id})
        item   = result.get("Item")
    except Exception as exc:
        return _error(500, str(exc))

    if not item:
        return _error(404, "Table not found")

    safe_cols = item.get("col_names", [])
    item["columns"] = safe_cols
    item.pop("col_names", None)

    if item.get("table_created") and item.get("pg_table_name"):
        # Read from PostgreSQL
        db_cfg = _get_superadmin_pg_connection()
        if db_cfg:
            try:
                conn     = _pg_connect(db_cfg)
                rows_raw = conn.run(f'SELECT * FROM "{item["pg_table_name"]}" LIMIT 1000')
                conn.close()
                item["rows"] = [dict(zip(safe_cols, r)) for r in rows_raw]
                return _ok({"table": item})
            except Exception as exc:
                logger.warning("PG read failed, falling back to S3: %s", exc)

    # Fallback: parse from S3
    s3_key = item.get("s3_key", "")
    if not s3_key:
        item["rows"] = []
        return _ok({"table": item})

    try:
        obj       = s3.get_object(Bucket=CSV_BUCKET, Key=s3_key)
        raw_bytes = obj["Body"].read()
        ext       = item.get("file_ext", "csv")
        _, rows   = _parse_file(raw_bytes, ext)
        item["rows"] = [dict(zip(safe_cols, r)) for r in rows[:1000]]
    except Exception as exc:
        logger.warning("S3 read failed: %s", exc)
        item["rows"] = []

    return _ok({"table": item})


def delete_table(event: dict, table_id: str) -> dict:
    """DELETE /csv/tables/{table_id} — delete S3 file, drop PG table if exists, remove DynamoDB record."""
    user, err = _get_user(event)
    if err:
        return err
    user_id = user.get("sub", "anonymous")

    try:
        result = meta_tbl.get_item(Key={"user_id": user_id, "table_id": table_id})
        item   = result.get("Item")
    except Exception as exc:
        return _error(500, str(exc))

    if not item:
        return _error(404, "Table not found")

    # Delete from S3
    s3_key = item.get("s3_key", "")
    if s3_key:
        try:
            s3.delete_object(Bucket=CSV_BUCKET, Key=s3_key)
        except Exception as exc:
            logger.warning("S3 delete failed for %s: %s", s3_key, exc)

    # Drop PG table if it was created
    if item.get("table_created") and item.get("pg_table_name"):
        db_cfg = _get_superadmin_pg_connection()
        if db_cfg:
            try:
                conn = _pg_connect(db_cfg)
                conn.run(f'DROP TABLE IF EXISTS "{item["pg_table_name"]}"')
                conn.close()
            except Exception as exc:
                logger.warning("PG DROP TABLE failed: %s", exc)

    # Remove DynamoDB metadata
    try:
        meta_tbl.delete_item(Key={"user_id": user_id, "table_id": table_id})
    except Exception as exc:
        return _error(500, str(exc))

    return _ok({"success": True})


def drop_table(event: dict, table_id: str) -> dict:
    """DELETE /csv/tables/{table_id}/drop-table — drop PG table only, keep S3 file and metadata."""
    user, err = _get_user(event)
    if err:
        return err
    user_id = user.get("sub", "anonymous")

    try:
        result = meta_tbl.get_item(Key={"user_id": user_id, "table_id": table_id})
        item   = result.get("Item")
    except Exception as exc:
        return _error(500, str(exc))

    if not item:
        return _error(404, "Table not found")

    if not item.get("table_created") or not item.get("pg_table_name"):
        return _error(400, "No PostgreSQL table has been created for this file")

    pg_table_name = item["pg_table_name"]

    # Drop PG table
    db_cfg = _get_superadmin_pg_connection()
    if not db_cfg:
        return _error(503, "No superadmin PostgreSQL connection found")
    try:
        conn = _pg_connect(db_cfg)
        conn.run(f'DROP TABLE IF EXISTS "{pg_table_name}"')
        conn.close()
    except Exception as exc:
        return _error(500, f"Failed to drop table: {exc}")

    # Reset metadata in DynamoDB (keep file, just clear table fields)
    try:
        meta_tbl.update_item(
            Key={"user_id": user_id, "table_id": table_id},
            UpdateExpression="SET table_created = :f REMOVE pg_table_name, connection_id, connection_name",
            ExpressionAttributeValues={":f": False},
        )
    except Exception as exc:
        return _error(500, str(exc))

    return _ok({"success": True, "dropped_table": pg_table_name})


# ── Lambda entry point ────────────────────────────────────────

def lambda_handler(event: dict, context) -> dict:
    method = (event.get("httpMethod") or "").upper()
    path   = event.get("path", "")

    if method == "OPTIONS":
        return {"statusCode": 200, "headers": CORS_HEADERS, "body": ""}

    try:
        if path == "/csv/upload" and method == "POST":
            return upload_file(event)

        if path == "/csv/tables" and method == "GET":
            return list_tables(event)

        if path.startswith("/csv/tables/"):
            remainder = path[len("/csv/tables/"):].strip("/")
            parts     = remainder.split("/")
            table_id  = parts[0]

            if len(parts) == 2 and parts[1] == "create-table" and method == "POST":
                return create_table(event, table_id)

            if len(parts) == 2 and parts[1] == "drop-table" and method == "DELETE":
                return drop_table(event, table_id)

            if len(parts) == 1:
                if method == "GET":
                    return get_table(event, table_id)
                if method == "DELETE":
                    return delete_table(event, table_id)

        return _error(404, f"Not found: {method} {path}")
    except Exception as exc:
        logger.exception("Unhandled error")
        return _error(500, str(exc))
